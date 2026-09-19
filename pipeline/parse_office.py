"""
Office / PDF attachment parsing
===============================
Three real shapes, each with a trap:
  XLSX  two-column key/value sheet
  DOCX  a table whose labels carry Chinese annotations - "Shipper (Principal or Seller) (发货人)"
  PDF   label and value on one line but **no colon** - "Shipper APRIL FINE PAPER TRADING", address continuation lines

Strategy: normalise everything to "Label: Value" text first, then reuse the single
resolve_label / parsing path. One field-parsing path, one alias table to maintain.
"""
from __future__ import annotations
import re, warnings, pathlib
warnings.filterwarnings("ignore")

from shipdoc_core.fields import FIELDS

# Strip CJK annotations: "Shipper (Principal or Seller) (发货人)" -> "Shipper (Principal or Seller)"
CJK_PAREN = re.compile(r"[（(]\s*[\u4e00-\u9fff][^)）]*[)）]")

# PDF has no colon: prefix-match known labels at line start. Longest first, so "Shipper" does not steal "Shipper/Exporter"
_ALL_LABELS: list[str] = sorted(
    {a for f in FIELDS for a in f.aliases} | {n for f in FIELDS for n in f.near_miss_labels},
    key=len, reverse=True,
)
_LABEL_PREFIX = re.compile(
    r"^\s*(" + "|".join(re.escape(l).replace(r"\ ", r"\s+") for l in _ALL_LABELS) + r")\s*[:\-]?\s+(.*)$",
    re.I,
)


CJK_ANY = re.compile(r"[\u3000-\u9fff\uff00-\uffef]+")


def _clean_label(s: str) -> str:
    """Strip CJK annotations first, then inline CJK. The real data has both."""
    s = CJK_PAREN.sub("", str(s))
    s = CJK_ANY.sub("", s)
    return re.sub(r"\s+", " ", s).strip().rstrip(":").strip()


def xlsx_to_kv_text(path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [c for c in row if c is not None and str(c).strip()]
            if len(cells) >= 2:
                lines.append(f"{_clean_label(cells[0])}: {cells[1]}")
            elif len(cells) == 1:
                lines.append(str(cells[0]))
    return "\n".join(lines)


def docx_to_kv_text(path) -> str:
    import docx
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if len(cells) >= 2:
                # line breaks inside a cell are address continuation; the first line is the entity
                value = cells[1].splitlines()[0].strip()
                lines.append(f"{_clean_label(cells[0])}: {value}")
            elif len(cells) == 1:
                lines.append(cells[0].splitlines()[0])
    return "\n".join(lines)


def pdf_to_kv_text(path) -> str:
    import pdfplumber
    out = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.splitlines():
                m = _LABEL_PREFIX.match(line)
                if m:
                    out.append(f"{_clean_label(m.group(1))}: {m.group(2).strip()}")
                else:
                    out.append(line)        # keep the raw line for document-type fingerprinting
    return "\n".join(out)


def office_to_text(path) -> str | None:
    """Convert pdf/docx/xlsx to 'Label: Value' text. Returns None on failure -> reported as unreadable."""
    p = pathlib.Path(path)
    try:
        ext = p.suffix.lower()
        if ext == ".xlsx":
            return xlsx_to_kv_text(p)
        if ext == ".docx":
            return docx_to_kv_text(p)
        if ext == ".pdf":
            txt = pdf_to_kv_text(p)
            # image-only scan: pdfplumber extracts no text -> reported as unreadable (no OCR in this system)
            return txt if txt and len(txt.strip()) > 40 else None
        return None
    except Exception:
        return None
