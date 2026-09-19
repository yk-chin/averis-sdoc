"""
Perturbation test - is the 1.0 generalisation or fitting?
=========================================================
Perturb a COPY of the dataset without changing its meaning, re-run the pipeline, score via /submit.
Whichever perturbation drops the score is the real weak point.

Usage:   python scripts_perturb.py            # run all
         python scripts_perturb.py P3 P5      # run selected
Output:  evals/perturbation.json (score + coverage per perturbation)

Hard constraints:
  - the original data/ is read-only; copies go to .cache/perturb/<P>/
  - ground_truth is never read; only inbox / attachments and the aggregate scores returned by /submit
  - evals/history.jsonl is not written (that is the mainline progress metric; perturbation runs do not count)
  - PDF attachments (28) are left untouched; every perturbation reports its actual coverage

Perturbations:
  P1  field labels replaced by synonyms (Port of Loading -> Load Port -> POL ...), different alias on SI and BL
  P2  company-suffix spelling (CO., LTD <-> Co Ltd <-> Company Limited), variant A on SI, variant B on BL
  P3  weight units: SI converted to MT (2 decimals), BL to LBS (1 decimal)
  P4a ports: SI keeps the name only (UN/LOCODE removed), BL untouched
  P4b ports: BL keeps the UN/LOCODE only - but only where name and code agree in the real LOCODE table
      (so the organiser's planted defects are not erased)
  P5  attachment names lose their _SI / _BL tags (email_001_SI.txt -> email_001_doc_a.txt)
"""
from __future__ import annotations
import json, pathlib, re, shutil, subprocess, sys, urllib.request

ROOT = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
from shipdoc_core.fields import FIELDS, FIELD_BY_KEY, FieldKind, resolve_label
from pipeline.parse_doc import LABEL_LINE

DATA = ROOT / "data"
WORK = ROOT / ".cache" / "perturb"
SERVER = "http://localhost:8080"
PY = sys.executable

# Real UN/LOCODEs (industry knowledge, not derived from the dataset). P4b turns a BL port into a bare code only when (code, name) agree.
REAL_LOCODE = {
    "MYPKG": "PORT KLANG", "SGSIN": "SINGAPORE", "INNSA": "NHAVA SHEVA", "VNSGN": "HO CHI MINH CITY",
    "USNYC": "NEW YORK", "KRPUS": "BUSAN", "CNSHA": "SHANGHAI", "CNNTG": "NANTONG", "PECLL": "CALLAO",
    "PKKHI": "KARACHI", "TRMER": "MERSIN", "KEMBA": "MOMBASA", "AUFRE": "FREMANTLE", "GNCKY": "CONAKRY",
    "LTKLJ": "KLAIPEDA", "NGAPP": "APAPA", "USSAV": "SAVANNAH", "KRPTK": "PYEONGTAEK", "MMRGN": "YANGON",
    "USBAL": "BALTIMORE", "AUBNE": "BRISBANE", "USLGB": "LONG BEACH", "ILASH": "ASHDOD", "PLGDN": "GDANSK",
    "CLVAP": "VALPARAISO", "SIKOP": "KOPER", "JOAQB": "AQABA", "USHOU": "HOUSTON", "PHCEB": "CEBU",
    "INTUT": "TUTICORIN",          # IDBUA (Buatan) could not be confirmed as an official code; left out
}


# ----------------------------------------------------------------------------
# Per-line transforms: (label, value, side, seed) -> (label, value). side in {"SI","BL"}
# ----------------------------------------------------------------------------
def _title(s: str) -> str:
    return " ".join(w.capitalize() if w.isalpha() else w for w in s.split())


def p1_label(label, value, side, seed):
    key, near = resolve_label(label)
    if key is None or near:
        return label, value
    aliases = [a for a in FIELD_BY_KEY[key].aliases if a != label.lower()]
    # different alias on SI and BL; the seed rotates it per email
    pick = aliases[(seed + (0 if side == "SI" else 3)) % len(aliases)]
    return _title(pick), value


_SUFFIX_A = [(r"\bCO\.?,?\s*LTD\.?", "Co Ltd"), (r"\bPTE\.?\s*LTD\.?", "Pte Ltd"),
             (r"\bSDN\.?\s*BHD\.?", "Sdn. Bhd."), (r"\bLLC\b", "L.L.C."), (r"\bGMBH\b", "GmbH"),
             (r"\bPTY\.?\s*LTD\.?", "Pty. Ltd.")]
_SUFFIX_B = [(r"\bCO\.?,?\s*LTD\.?", "Company Limited"), (r"\bPTE\.?\s*LTD\.?", "Pte. Limited"),
             (r"\bSDN\.?\s*BHD\.?", "Sendirian Berhad"), (r"\bLLC\b", "LLC"), (r"\bGMBH\b", "GmbH"),
             (r"\bPTY\.?\s*LTD\.?", "Pty Limited")]


def p2_suffix(label, value, side, seed):
    key, near = resolve_label(label)
    if key is None or near or FIELD_BY_KEY[key].kind is not FieldKind.PARTY:
        return label, value
    table = _SUFFIX_A if side == "SI" else _SUFFIX_B
    v = value
    for pat, rep in table:
        v = re.sub(pat, rep, v, flags=re.I)
    return label, v


def p3_weight(label, value, side, seed):
    key, near = resolve_label(label)
    if key is None or near or FIELD_BY_KEY[key].kind is not FieldKind.WEIGHT_KG:
        return label, value
    m = re.search(r"\d[\d,]*(?:\.\d+)?", value)
    if not m:
        return label, value
    kg = float(m.group(0).replace(",", ""))
    if side == "SI":
        mt = kg / 1000
        new = f"{mt:g} MT" if mt == int(mt) else f"{mt:.2f} MT"      # 2 decimals avoid the thousands-separator ambiguity
    else:
        new = f"{kg * 2.20462262:,.1f} LBS"
    return label, new


def p4a_port_si_name_only(label, value, side, seed):
    key, near = resolve_label(label)
    if key is None or near or FIELD_BY_KEY[key].kind is not FieldKind.PORT or side != "SI":
        return label, value
    return label, re.sub(r"\s*\([A-Z]{2}[A-Z0-9]{3}\)\s*", "", value).strip()


def p4b_port_bl_code_only(label, value, side, seed):
    key, near = resolve_label(label)
    if key is None or near or FIELD_BY_KEY[key].kind is not FieldKind.PORT or side != "BL":
        return label, value
    m = re.search(r"\(([A-Z]{2}[A-Z0-9]{3})\)", value)
    if not m:
        return label, value
    code = m.group(1)
    name = re.sub(r"\s*\(.*?\)\s*", " ", value).split(",")[0].strip().upper()
    if REAL_LOCODE.get(code) and REAL_LOCODE[code] in name:      # only replace when name and code agree, keeping planted defects intact
        return label, code
    return label, value


PERTURBATIONS = {
    "P1": ("field-label synonyms rotated", p1_label),
    "P2": ("company-suffix spelling (SI variant A / BL variant B)", p2_suffix),
    "P3": ("weight units (SI->MT / BL->LBS)", p3_weight),
    "P4a": ("ports: SI name only", p4a_port_si_name_only),
    "P4b": ("ports: BL LOCODE only (consistent pairs)", p4b_port_bl_code_only),
    "P5": ("attachment names without _SI/_BL", None),
}


# ----------------------------------------------------------------------------
# apply a transform to one attachment
# ----------------------------------------------------------------------------
def _side(name: str) -> str:
    return "SI" if "_SI." in name.upper() else "BL"


def _seed(name: str) -> int:
    m = re.search(r"email_(\d+)", name)
    return int(m.group(1)) if m else 0


def apply_txt(path: pathlib.Path, fn) -> int:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    out, changed = [], 0
    for line in lines:
        m = LABEL_LINE.match(line)
        if m:
            lab, val = m.group(1).strip(), m.group(2).strip()
            nl, nv = fn(lab, val, _side(path.name), _seed(path.name))
            if (nl, nv) != (lab, val):
                changed += 1
                line = f"{nl}: {nv}"
        out.append(line)
    path.write_text("\n".join(out) + "\n", encoding="utf-8")
    return changed


def apply_xlsx(path: pathlib.Path, fn) -> int:
    import openpyxl
    wb = openpyxl.load_workbook(path)
    changed = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            cells = [c for c in row if c.value is not None and str(c.value).strip()]
            if len(cells) >= 2:
                lab, val = str(cells[0].value).strip().rstrip(":"), str(cells[1].value).strip()
                nl, nv = fn(lab, val, _side(path.name), _seed(path.name))
                if (nl, nv) != (lab, val):
                    cells[0].value, cells[1].value = nl, nv; changed += 1
    wb.save(path)
    return changed


def apply_docx(path: pathlib.Path, fn) -> int:
    import docx
    d = docx.Document(str(path))
    changed = 0
    for p in d.paragraphs:
        m = LABEL_LINE.match(p.text)
        if m:
            nl, nv = fn(m.group(1).strip(), m.group(2).strip(), _side(path.name), _seed(path.name))
            if (nl, nv) != (m.group(1).strip(), m.group(2).strip()):
                p.text = f"{nl}: {nv}"; changed += 1
    for t in d.tables:
        for row in t.rows:
            cells = [c for c in row.cells if c.text.strip()]
            if len(cells) >= 2:
                lab = cells[0].text.strip().rstrip(":")
                first, *rest = cells[1].text.split("\n")
                nl, nv = fn(lab, first.strip(), _side(path.name), _seed(path.name))
                if (nl, nv) != (lab, first.strip()):
                    cells[0].text = nl; cells[1].text = "\n".join([nv, *rest]); changed += 1
    d.save(str(path))
    return changed


def build_copy(name: str, fn) -> dict:
    dst = WORK / name
    if dst.exists():
        shutil.rmtree(dst)
    shutil.copytree(DATA, dst)
    cov = {"txt": [0, 0], "xlsx": [0, 0], "docx": [0, 0], "pdf": [0, 0], "lines_changed": 0}
    if fn is None:                                             # P5: rename files + inbox references
        renamed = 0
        for p in sorted((dst / "attachments").glob("*")):
            new = p.name.replace("_SI.", "_doc_a.").replace("_BL.", "_doc_b.")
            if new != p.name:
                p.rename(p.with_name(new)); renamed += 1
        for f in (dst / "inbox").glob("*.json"):
            e = json.load(open(f, encoding="utf-8"))
            e["attachments"] = [a.replace("_SI.", "_doc_a.").replace("_BL.", "_doc_b.") for a in e.get("attachments") or []]
            json.dump(e, open(f, "w", encoding="utf-8"), indent=2, ensure_ascii=False)
        cov["renamed"] = renamed
        return cov
    for p in sorted((dst / "attachments").glob("*")):
        ext = p.suffix.lstrip(".").lower()
        cov[ext][1] += 1
        if ext == "txt":
            n = apply_txt(p, fn)
        elif ext == "xlsx":
            n = apply_xlsx(p, fn)
        elif ext == "docx":
            n = apply_docx(p, fn)
        else:
            continue                                            # pdf untouched
        if n:
            cov[ext][0] += 1; cov["lines_changed"] += n
    return cov


# ----------------------------------------------------------------------------
def run_and_score(data_dir: pathlib.Path, out: pathlib.Path) -> dict:
    r = subprocess.run([PY, str(ROOT / "pipeline" / "run.py"), str(data_dir), str(out)],
                       capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        raise RuntimeError(r.stderr[-800:])
    sub = json.load(open(out, encoding="utf-8"))
    req = urllib.request.Request(f"{SERVER}/submit", data=json.dumps(sub).encode(),
                                 headers={"Content-Type": "application/json"}, method="POST")
    s = json.loads(urllib.request.urlopen(req, timeout=120).read())
    return {"final_score": s["final_score"], "macro_f1": s["stage1"]["macro_f1"],
            "defect_f1": s["stage3"]["defect_f1"], "field_f1": s["stage3"]["field_f1"],
            "end_to_end": s["end_to_end"]["rate"],
            "esc_precision": s["reliability"]["escalation_precision"],
            "esc_recall": s["reliability"]["escalation_recall"],
            "pred_review": s["reliability"]["pred_review"]}


def main(selected: list[str]):
    WORK.mkdir(parents=True, exist_ok=True)
    print("baseline (unperturbed) ...")
    base = run_and_score(DATA, WORK / "baseline_submission.json")
    results = {"baseline": base}
    rows = [("baseline", "unperturbed", base, {})]
    for name, (desc, fn) in PERTURBATIONS.items():
        if selected and name not in selected:
            continue
        print(f"{name} {desc} …", flush=True)
        cov = build_copy(name, fn)
        try:
            sc = run_and_score(WORK / name, WORK / f"{name}_submission.json")
        except RuntimeError as e:
            sc = {"error": str(e)}
        results[name] = {"desc": desc, "coverage": cov, **sc}
        rows.append((name, desc, sc, cov))

    print("\n| perturbation | description | final | delta | macro_f1 | defect_f1 | field_f1 | end_to_end | esc_P | esc_R | coverage |")
    print("|---|---|---|---|---|---|---|---|---|---|---|")
    for name, desc, sc, cov in rows:
        if "error" in sc:
            print(f"| {name} | {desc} | ERROR | | | | | | | | {sc['error'][:60]} |"); continue
        d = sc["final_score"] - base["final_score"]
        if cov.get("renamed") is not None:
            c = f"{cov['renamed']} files renamed"
        elif cov:
            c = " ".join(f"{k}:{v[0]}/{v[1]}" for k, v in cov.items() if isinstance(v, list) and v[1])
            c += f" ({cov['lines_changed']} lines)"
        else:
            c = "—"
        print(f"| {name} | {desc} | {sc['final_score']:.4f} | {d:+.4f} | {sc['macro_f1']:.4f} | {sc['defect_f1']:.4f} | "
              f"{sc['field_f1']:.3f} | {sc['end_to_end']:.4f} | {sc['esc_precision']:.3f} | {sc['esc_recall']:.3f} | {c} |")
    (ROOT / "evals").mkdir(exist_ok=True)
    json.dump(results, open(ROOT / "evals" / "perturbation.json", "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print(f"\nWritten to evals/perturbation.json; copies under {WORK}")


if __name__ == "__main__":
    main(sys.argv[1:])
